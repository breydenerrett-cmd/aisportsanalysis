"""src/hounds/bet_export.py: the bets spreadsheet (owner ask 2026-10-28).

What these pin: every bet is a row; the hound's own recorded reason is in it; a replay row is labelled and never counted; a missing
number is an empty cell, never a zero; the review columns are labelled hindsight and filled only from a real review; the CSV opens
in Excel (byte-order mark) and can never run a formula; the bet id is the same value the record page uses.
"""
from __future__ import annotations

import csv
import io
import unittest
from datetime import date, datetime, timezone

from src.ai import postmortem_store as store
from src.hounds import bet_export as bx
from src.hounds import contract, stats, track_record
from src.hounds.ledgers import Bet
from tests.postmortem_fixtures import bet, hound_data

NOW = datetime(2026, 10, 28, 23, 0, tzinfo=timezone.utc)


def rows_of(data, **kw):
    return bx.build_rows(data, now=NOW, **kw)


def parse_csv(raw: bytes):
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))


def review(**over):
    out = {"bet_key": "x", "depth": "fast", "model": "claude-haiku-5-5", "created_utc": "2026-10-28T21:00:00Z", "label": "Hindsight review",
           "what_happened": "Lost 5-2 after the starter left early.", "key_factors": ["starter left", "bullpen"], "factor_tags": ["starter_performance"],
           "reason_relevance": "partly_valid", "reason_assessment": "Fine angle, hurt starter.", "luck_vs_process": "mixed",
           "info_hound_lacked": ["a reported blister"], "lesson": "Check the health report.", "grounded": True, "confidence": "high",
           "analyst_takes": [{"outlet": "Example", "take": "Thin lineup.", "url": "https://example.com/recap"}],
           "sources": [{"url": "https://example.com/recap", "title": "Recap", "kind": "recap"}]}
    out.update(over)
    return out


class TheColumns(unittest.TestCase):
    def test_every_bet_is_a_row_with_every_column_the_owner_asked_for(self):
        data = hound_data("u-1", [bet(1, result="loss"), bet(2, result="win", price=150)])
        rows = rows_of(data)
        self.assertEqual(len(rows), 2)
        for needed in ("Slate date", "Sport", "League", "Matchup", "Market", "Side", "Pick", "Line", "Price taken", "Average price (board)",
                       "Closing-line value (pts)", "Stake (paper units)", "Hound's recorded reason (before the game)", "Segment", "Result",
                       "Units won or lost", "Hindsight: what happened", "Hindsight: lesson", "Hindsight: sources"):
            self.assertIn(needed, bx.HEADERS)
        self.assertEqual([k for k, _h in bx.COLUMNS], [k for k, _h in bx.COLUMNS if rows[0].keys() >= {k}])    # every column has a cell
        r = rows[0]
        self.assertEqual((r["date"], r["matchup"], r["market"], r["pick"], r["price"], r["price_average"], r["clv_pts"]),
                         ("2026-10-28", "Team A1 at Team B1", "h2h", "Team B1 moneyline", 120, 115, 0.6))
        self.assertEqual((r["sport"], r["league"]), ("Baseball", "MLB"))
        self.assertEqual((r["result"], r["units"]), ("loss", -1.0))

    def test_the_hounds_recorded_reason_is_in_the_row_as_stored(self):
        rows = rows_of(hound_data("u-1", [bet(1, reason="Rested starter, plus money, under 7.5 total.")]))
        self.assertEqual(rows[0]["reason"], "Rested starter, plus money, under 7.5 total.")

    def test_replay_is_labelled_and_never_counted_and_an_open_ticket_is_not_yet(self):
        data = hound_data("u-1", [bet(1, result="win", price=120), bet(2, result="loss", segment="replay"), bet(3, result=None), bet(4, result="void"),
                                  bet(5, result="push")])
        by_game = {r["matchup"]: r for r in rows_of(data)}
        fwd, replay, open_, void, push = (by_game[f"Team A{n} at Team B{n}"] for n in (1, 2, 3, 4, 5))
        self.assertEqual((fwd["segment"], fwd["counted"]), ("forward", "yes"))
        self.assertEqual(replay["segment"], "replay (hindsight, never counted)")
        self.assertEqual(replay["counted"], "no")
        self.assertEqual((open_["result"], open_["counted"], open_["status"]), ("open", "not settled yet", "OPEN"))
        self.assertEqual(void["counted"], "no")
        self.assertEqual(push["counted"], "yes")                       # a push is in the record (as a push)

    def test_a_missing_clv_or_price_is_an_empty_cell_not_a_zero(self):
        r = rows_of(hound_data("u-1", [bet(1, clv_pts=None, price_average=None)]))[0]
        self.assertIsNone(r["clv_pts"])
        self.assertIsNone(r["price_average"])
        row = parse_csv(bx.to_csv([r]))[1]
        self.assertEqual(row[bx.HEADERS.index("Closing-line value (pts)")], "")
        self.assertEqual(row[bx.HEADERS.index("Average price (board)")], "")

    def test_stake_is_the_members_own_stake_and_one_unit_for_the_house(self):
        member = rows_of(hound_data("u-1", [bet(1, stake_units=2.0)], kind="user"))[0]
        house = rows_of(hound_data("h-1", [bet(1, hound_id="h-1", stake_units=2.0)], kind="house"))[0]
        plain = rows_of(hound_data("u-1", [bet(1)], kind="user"))[0]
        self.assertEqual((member["stake_units"], house["stake_units"], plain["stake_units"]), (2.0, 1.0, 1.0))

    def test_the_stake_reaches_the_row_from_the_ledger_bet(self):
        row = stats.bet_row(Bet(hound_id="u-1", date="2026-10-28", stake_units=2.5), None)
        self.assertEqual(row.stake_units, 2.5)
        self.assertIsNone(contract.BetRow(hound_id="x", date="2026-10-28").stake_units)

    def test_side_is_inferred_from_the_pick_and_the_game(self):
        r = rows_of(hound_data("u-1", [bet(1)]))[0]
        self.assertEqual(r["side"], "home")
        over = bet(2, market="totals", selection="Over 8.5", line="8.5")
        self.assertEqual(rows_of(hound_data("u-1", [over]))[0]["side"], "over")

    def test_the_bet_id_is_the_value_the_record_page_uses_for_the_same_bet(self):
        data = hound_data("u-1", [bet(1, result="loss")])
        sheet = rows_of(data)[0]["bet_id"]
        page = track_record.build(data, as_of="2026-10-28", now=NOW)["tickets"][0]["bet_key"]
        self.assertEqual(sheet, page)
        self.assertRegex(sheet, r"^pm_[0-9a-f]{16}$")

    def test_rows_are_oldest_first(self):
        data = hound_data("u-1", [bet(3, date="2026-10-28"), bet(1, date="2026-10-26"), bet(2, date="2026-10-27")])
        self.assertEqual([r["date"] for r in rows_of(data)], ["2026-10-26", "2026-10-27", "2026-10-28"])

    def test_an_empty_hound_gives_a_sheet_with_just_the_header(self):
        raw = bx.to_csv(rows_of(hound_data("u-1", [])))
        self.assertEqual(parse_csv(raw), [bx.HEADERS])


class TheHindsightColumns(unittest.TestCase):
    def setUp(self):
        self.data = hound_data("u-1", [bet(1, result="loss"), bet(2, result="loss"), bet(3, result="loss", segment="replay"), bet(4, result=None)])
        self.keys = [store.bet_key("u-1", b) for b in self.data.bets]

    def test_a_reviewed_bet_fills_every_hindsight_column_and_the_rest_say_why_not(self):
        rows = {r["bet_id"]: r for r in rows_of(self.data, reviews={self.keys[0]: review()}, queued={self.keys[1]})}
        done = rows[self.keys[0]]
        self.assertEqual(done["review_status"], "reviewed")
        for key in ("what_happened", "key_factors", "reason_relevance", "reason_assessment", "luck_vs_process", "info_lacked", "lesson",
                    "analyst_takes", "sources", "confidence", "review_model", "review_time", "review_depth"):
            self.assertTrue(done[key], key)
        self.assertIn("https://example.com/recap", done["sources"])
        self.assertIn("Example: Thin lineup.", done["analyst_takes"])
        self.assertEqual(rows[self.keys[1]]["review_status"], "queued")
        self.assertEqual(rows[self.keys[1]]["lesson"], "")
        self.assertEqual(rows[self.keys[2]]["review_status"], "n/a")                       # a replay bet is never reviewed
        self.assertEqual(rows[self.keys[3]]["review_status"], "not reviewed")

    def test_an_unsourced_review_says_so_in_the_status(self):
        rows = {r["bet_id"]: r for r in rows_of(self.data, reviews={self.keys[0]: review(grounded=False, sources=[], confidence="low")})}
        self.assertEqual(rows[self.keys[0]]["review_status"], "reviewed (no sources found)")

    def test_every_review_column_is_headed_hindsight(self):
        first = [i for i, (k, _h) in enumerate(bx.COLUMNS) if k == "review_status"][0]
        for _k, header in bx.COLUMNS[first:]:
            self.assertTrue(header.startswith("Hindsight:"), header)
        for _k, header in bx.COLUMNS[:first]:
            self.assertFalse(header.startswith("Hindsight:"), header)

    def test_no_ledger_at_all_still_gives_a_complete_sheet(self):
        rows = rows_of(self.data)
        self.assertEqual(len(rows), 4)
        self.assertEqual({r["review_status"] for r in rows}, {"not reviewed", "n/a"})


class TheCsv(unittest.TestCase):
    def test_it_starts_with_a_byte_order_mark_so_excel_reads_it_as_utf8(self):
        raw = bx.to_csv(rows_of(hound_data("u-1", [bet(1, game="Atlético at Münster")])))
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        self.assertIn("Atlético at Münster", parse_csv(raw)[1])

    def test_a_cell_that_would_run_as_a_formula_is_neutralised(self):
        evil = ['=HYPERLINK("http://evil.example","x")', "+SUM(1,1)", "-2+3", "@cmd", "\tx", "\rx"]
        for text in evil:
            self.assertTrue(str(bx.safe_cell(text)).startswith("'"), text)
        r = rows_of(hound_data("u-1", [bet(1, reason='=cmd|" /C calc"!A0')]))[0]
        r["lesson"] = "+1 lesson"
        cells = parse_csv(bx.to_csv([r]))[1]
        self.assertEqual(cells[bx.HEADERS.index("Hound's recorded reason (before the game)")], "'=cmd|\" /C calc\"!A0")
        self.assertEqual(cells[bx.HEADERS.index("Hindsight: lesson")], "'+1 lesson")

    def test_numbers_and_plain_numeric_text_are_left_alone(self):
        for ok in (-110, 120, 0.6, "-1.5", "+3", "8.5", ".5"):
            self.assertEqual(bx.safe_cell(ok), ok)
        r = rows_of(hound_data("u-1", [bet(1, price=-130, line="-1.5")]))[0]
        cells = parse_csv(bx.to_csv([r]))[1]
        self.assertEqual(cells[bx.HEADERS.index("Price taken")], "-130")
        self.assertEqual(cells[bx.HEADERS.index("Line")], "-1.5")

    def test_multi_line_review_text_and_commas_survive_a_round_trip(self):
        data = hound_data("u-1", [bet(1)])
        key = store.bet_key("u-1", data.bets[0])
        rows = rows_of(data, reviews={key: review(what_happened='He said "no", then left,\nlater.')})
        cells = parse_csv(bx.to_csv(rows))[1]
        self.assertEqual(cells[bx.HEADERS.index("Hindsight: what happened")], 'He said "no", then left,\nlater.')

    def test_every_row_has_exactly_as_many_cells_as_headers(self):
        raw = bx.to_csv(rows_of(hound_data("u-1", [bet(1), bet(2, segment="replay"), bet(3, result=None)])))
        for row in parse_csv(raw):
            self.assertEqual(len(row), len(bx.HEADERS))


class TheFormats(unittest.TestCase):
    def test_csv_is_always_offered_and_xlsx_only_when_openpyxl_is_installed(self):
        formats = bx.available_formats()
        self.assertIn("csv", formats)
        try:
            import openpyxl  # noqa: F401
            self.assertIn("xlsx", formats)
        except ImportError:
            self.assertEqual(formats, ("csv",))
            self.assertIsNone(bx.to_xlsx([]))

    def test_the_xlsx_round_trips_when_it_can_be_made(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed (it is not in api/requirements.txt)")
        raw = bx.to_xlsx(rows_of(hound_data("u-1", [bet(1)])))
        sheet = openpyxl.load_workbook(io.BytesIO(raw)).active
        self.assertEqual([c.value for c in sheet[1]], bx.HEADERS)
        self.assertEqual(sheet.max_row, 2)

    def test_the_file_name_is_safe_and_dated(self):
        self.assertEqual(bx.filename("u-1", "csv", date(2026, 10, 28)), "linehound-u-1-bets-2026-10-28.csv")
        name = bx.filename("../../etc/passwd; rm", "csv", date(2026, 10, 28))
        self.assertNotIn("/", name)
        self.assertNotIn(" ", name)
        self.assertNotIn(";", name)


class TheAllVersionsSheet(unittest.TestCase):
    """``versions=all``: every bet of every version, a Version and a Version status column, the default sheet untouched."""

    def versions(self):
        v1 = hound_data("u-1", [bet(1, hound_id="u-1", result="loss"), bet(2, hound_id="u-1", result="win", price=150)])
        v2 = hound_data("u-2", [bet(3, hound_id="u-2", result="win"), bet(4, hound_id="u-2", result=None), bet(5, hound_id="u-2", result="loss", segment="replay")])
        return [{"data": v1, "label": "v1", "status": "replaced by a later version", "reviews": {}, "queued": set()},
                {"data": v2, "label": "v2", "status": "running"}]

    def test_the_default_columns_and_rows_are_exactly_what_they_were(self):
        rows = rows_of(hound_data("u-1", [bet(1)]))
        self.assertNotIn("version", rows[0])
        self.assertNotIn("version_status", rows[0])
        self.assertEqual(parse_csv(bx.to_csv(rows))[0], bx.HEADERS)
        self.assertEqual(bx.COLUMNS_ALL_VERSIONS[:3], bx.COLUMNS[:3])
        self.assertEqual(bx.COLUMNS_ALL_VERSIONS[5:], bx.COLUMNS[3:])
        self.assertEqual([h for _k, h in bx.COLUMNS_ALL_VERSIONS[3:5]], ["Version", "Version status"])

    def test_every_bet_of_every_version_is_a_row_oldest_version_first(self):
        rows = bx.build_rows_all_versions(self.versions(), now=NOW)
        self.assertEqual(len(rows), 2 + 3)
        self.assertEqual([r["version"] for r in rows], ["v1", "v1", "v2", "v2", "v2"])
        self.assertEqual([r["version_status"] for r in rows], ["replaced by a later version"] * 2 + ["running"] * 3)
        self.assertEqual([r["hound_id"] for r in rows], ["u-1", "u-1", "u-2", "u-2", "u-2"])
        ids = [r["bet_id"] for r in rows]
        self.assertEqual(len(set(ids)), len(ids))

    def test_the_replay_and_open_rows_keep_their_labels_in_every_version(self):
        rows = bx.build_rows_all_versions(self.versions(), now=NOW)
        by = {(r["version"], r["counted"]): 0 for r in rows}
        self.assertIn(("v2", bx.COUNTED_OPEN), by)
        self.assertIn(("v2", bx.COUNTED_NO), by)
        replay = [r for r in rows if r["segment"] == bx.SEGMENT_REPLAY]
        self.assertEqual([(r["version"], r["review_status"]) for r in replay], [("v2", bx.REVIEW_NA)])

    def test_the_csv_has_the_two_extra_columns_and_every_cell_is_still_neutralised(self):
        rows = bx.build_rows_all_versions(self.versions(), now=NOW)
        rows[0]["version_status"] = "=HYPERLINK(\"http://x\")"
        sheet = parse_csv(bx.to_csv(rows, bx.COLUMNS_ALL_VERSIONS))
        head = sheet[0]
        self.assertEqual(head, [h for _k, h in bx.COLUMNS_ALL_VERSIONS])
        self.assertEqual(sheet[1][head.index("Version")], "v1")
        self.assertTrue(sheet[1][head.index("Version status")].startswith("'="))
        self.assertTrue(all(len(r) == len(head) for r in sheet))

    def test_the_xlsx_has_them_too_when_it_can_be_made(self):
        try:
            import openpyxl
        except ImportError:
            self.assertIsNone(bx.to_xlsx([], bx.COLUMNS_ALL_VERSIONS))
            return
        raw = bx.to_xlsx(bx.build_rows_all_versions(self.versions(), now=NOW), bx.COLUMNS_ALL_VERSIONS)
        sheet = openpyxl.load_workbook(io.BytesIO(raw)).active
        self.assertEqual([c.value for c in sheet[1]], [h for _k, h in bx.COLUMNS_ALL_VERSIONS])
        self.assertEqual(sheet.max_row, 6)

    def test_the_reviews_of_each_version_land_on_its_own_rows(self):
        versions = self.versions()
        key = store.bet_key("u-1", versions[0]["data"].bets[0])
        versions[0]["reviews"] = {key: review(grounded=True)}
        rows = bx.build_rows_all_versions(versions, now=NOW)
        mine = next(r for r in rows if r["bet_id"] == key)
        self.assertEqual((mine["version"], mine["review_status"]), ("v1", "reviewed"))
        self.assertEqual(sum(1 for r in rows if r["review_status"] == "reviewed"), 1)

    def test_the_file_name_says_all_versions_only_when_it_is_the_all_versions_sheet(self):
        self.assertEqual(bx.filename("u-1", "csv", date(2026, 10, 28)), "linehound-u-1-bets-2026-10-28.csv")
        self.assertEqual(bx.filename("u-1", "csv", date(2026, 10, 28), all_versions=True), "linehound-u-1-bets-all-versions-2026-10-28.csv")

    def test_a_lineage_with_no_versions_is_an_empty_sheet_with_headers(self):
        self.assertEqual(parse_csv(bx.to_csv(bx.build_rows_all_versions([], now=NOW), bx.COLUMNS_ALL_VERSIONS)),
                         [[h for _k, h in bx.COLUMNS_ALL_VERSIONS]])


class TheCopy(unittest.TestCase):
    def test_no_banned_word_in_a_column_header_or_a_label(self):
        import re
        banned = re.compile(r"(?<![-_A-Za-z])dogs?(?![-_A-Za-z])|\bbeta\b|\bnew\b|nothing cleared the bar|\bprofit|guarantee", re.I)
        self.assertIsNone(banned.search(" ".join(bx.HEADERS)))
        self.assertIsNone(banned.search(" ".join(h for _k, h in bx.COLUMNS_ALL_VERSIONS)))
        self.assertIsNone(banned.search("replaced by a later version running retired buried benched"))
        self.assertIsNone(banned.search(bx.SEGMENT_REPLAY + bx.COUNTED_OPEN))


if __name__ == "__main__":
    unittest.main()
